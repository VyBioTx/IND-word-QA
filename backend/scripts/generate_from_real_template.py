#!/usr/bin/env -S pixi run python
"""Create a test docx by modifying the real report as template, embedding all 14 issue types.

Works at the OPC/XML level — does NOT use python-docx to re-save the original document,
preserving the original structure (OLE objects, charts, custom XML, etc.).
"""

from __future__ import annotations
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from lxml import etree
import re
import shutil

REAL_REPORT = Path("/tmp/opencode/study_report.docx")
OUTPUT_DIR = Path("/tmp/opencode")
OUTPUT = OUTPUT_DIR / "study_report_modified.docx"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSMAP = {"w": W}


def _make_p(text: str) -> etree.Element:
    p = etree.SubElement(etree.Element(f"{{{W}}}p", nsmap=NSMAP), f"{{{W}}}r")
    t = etree.SubElement(p, f"{{{W}}}t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    return p.getparent()





def _make_table(headers: list[str], rows: list[list[str]], col_widths: list[str] | None = None) -> etree.Element:
    tbl = etree.Element(f"{{{W}}}tbl", nsmap=NSMAP)
    tblPr = etree.SubElement(tbl, f"{{{W}}}tblPr")
    tblW = etree.SubElement(tblPr, f"{{{W}}}tblW")
    tblW.set(f"{{{W}}}w", "5000")
    tblW.set(f"{{{W}}}type", "pct")

    tblGrid = etree.SubElement(tbl, f"{{{W}}}tblGrid")
    for i in range(len(headers)):
        gridCol = etree.SubElement(tblGrid, f"{{{W}}}gridCol")
        gridCol.set(f"{{{W}}}w", col_widths[i] if col_widths else "1000")

    def _add_row(cells: list[str]) -> None:
        tr = etree.SubElement(tbl, f"{{{W}}}tr")
        for cell_text in cells:
            tc = etree.SubElement(tr, f"{{{W}}}tc")
            tcPr = etree.SubElement(tc, f"{{{W}}}tcPr")
            tcW = etree.SubElement(tcPr, f"{{{W}}}tcW")
            tcW.set(f"{{{W}}}w", "1000")
            tcW.set(f"{{{W}}}type", "dxa")
            p = etree.SubElement(tc, f"{{{W}}}p")
            r = etree.SubElement(p, f"{{{W}}}r")
            t = etree.SubElement(r, f"{{{W}}}t")
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            t.text = cell_text

    _add_row(headers)
    for row in rows:
        _add_row(row)
    return tbl


def copy_and_inject(path: Path = OUTPUT) -> Path:
    shutil.copy(str(REAL_REPORT), str(path))

    with ZipFile(str(path), "r") as source:
        members = {name: source.read(name) for name in source.namelist()}

    # ---- Parse document.xml ----
    doc_xml = members["word/document.xml"].decode("utf-8")
    doc_root = etree.fromstring(doc_xml.encode("utf-8"))
    body = doc_root.find(f"{{{W}}}body")

    # Save sectPr (must be absolute last element in body), then remove it
    sectPrs = body.findall(f"{{{W}}}sectPr")
    for sp in sectPrs:
        body.remove(sp)

    # ---- TEXT-001: Duplicate word ----
    body.append(_make_p("The the study demonstrated significant antitumor activity."))
    # ---- TEXT-002: Repeated spaces ----
    body.append(_make_p("The  efficacy  was evaluated  on day 21."))
    # ---- NUM-001: Wrong percentage ----
    body.append(_make_p("5/8 animals (75.0%) showed tumor regression in the high dose group."))
    # ---- TABLE-001: Wrong table total ----
    body.append(_make_table(["Group", "Count"], [["Vehicle", "8"], ["Low dose", "8"], ["High dose", "8"], ["Total", "22"]]))
    # ---- TABLE-002: Empty cell ----
    body.append(_make_table(["Parameter", "Value"], [["Cmax", ""], ["Tmax", "2h"]]))
    # ---- TABLE-003: N/A inconsistency ----
    body.append(_make_table(["Status"], [["Not Applicable"], ["N/A"], ["NA"]]))
    # ---- REF-001: Missing table/figure reference ----
    body.append(_make_p("As shown in Figure 10, the tumor growth was inhibited."))
    body.append(_make_p("Detailed PK parameters are listed in Table 25."))
    # ---- HEAD-001: Heading numbering gap ----
    body.append(_make_p("6.1  Summary of Findings"))
    body.append(_make_p("6.3  Future Directions"))
    # ---- ABBR-001: Undefined abbreviation ----
    body.append(_make_p("MTD was determined based on the toxicity profile observed during the study."))
    body.append(_make_p("SAE were reported and followed up according to the protocol."))
    # ---- ABBR-002: Multiple definitions ----
    body.append(_make_p("The maximum tolerated dose (MTD) was established at 30 mg/kg."))
    body.append(_make_p("Mean tumor diameter (MTD) was also recorded at each time point."))
    # ---- TERM-001: Terminology inconsistency ----
    body.append(_make_p("PROTOCOL-001 was followed for all procedures."))
    body.append(_make_p("All deviations from PROTOCOL001 were documented."))

    # Re-add sectPr as the last element in body
    for sp in sectPrs:
        body.append(sp)

    members["word/document.xml"] = etree.tostring(doc_root, xml_declaration=True, encoding="UTF-8", standalone=True)

    rels_path = "word/_rels/document.xml.rels"

    # ---- WORD-001: Header/footer version mismatch ----
    # Put TWO version strings in footer so the header/footer rule catches them
    for fname in ["word/footer1.xml"]:
        if fname in members:
            root = etree.fromstring(members[fname])
            root.insert(0, _make_p("Protocol v2.0 / Confidential - Version 2.1"))
            members[fname] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ---- WORD-002: Insert a Word comment ----
    members["word/comments.xml"] = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:comment w:id="0" w:author="Reviewer" w:date="2022-08-01T10:00:00Z">
    <w:p><w:r><w:t>Please verify the statistical analysis for Group 5.</w:t></w:r></w:p>
  </w:comment>
</w:comments>""".encode("utf-8")

    # Register comments.xml in [Content_Types].xml
    ct_xml = members["[Content_Types].xml"].decode("utf-8")
    if "word/comments.xml" not in ct_xml:
        ct_xml = ct_xml.replace("</Types>", '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/></Types>', 1)
    members["[Content_Types].xml"] = ct_xml.encode("utf-8")

    # Add comments relationship in document.xml.rels
    rels_xml = members[rels_path].decode("utf-8")
    if "comments" not in rels_xml:
        existing_ids = re.findall(r'Id="rId(\d+)"', rels_xml)
        next_id = max(int(x) for x in existing_ids) + 1 if existing_ids else 1
        rels_xml = rels_xml.replace("</Relationships>", f'<Relationship Id="rId{next_id}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/></Relationships>', 1)
    members[rels_path] = rels_xml.encode("utf-8")

    # ---- WORD-003: Inject a tracked change with unique w:id ----
    doc_xml2 = members["word/document.xml"].decode("utf-8")
    existing_ids = [int(x) for x in re.findall(r'w:id="(\d+)"', doc_xml2) if x.isdigit()]
    next_id = max(existing_ids) + 1 if existing_ids else 2
    insert_marker = f'<w:ins w:id="{next_id}" w:author="Author" w:date="2022-07-30T14:00:00Z"><w:r><w:t>revised value</w:t></w:r></w:ins>'
    doc_xml2 = doc_xml2.replace("<w:body>", f"<w:body>{insert_marker}", 1)
    members["word/document.xml"] = doc_xml2.encode("utf-8")

    with ZipFile(str(path), "w", ZIP_DEFLATED) as target:
        for name, content in members.items():
            target.writestr(name, content)
    return path


if __name__ == "__main__":
    path = copy_and_inject()
    print(f"Created: {path}")
    # Verify
    from docx import Document
    doc = Document(str(path))
    print(f"Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
